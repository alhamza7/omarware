# -*- coding: utf-8 -*-
"""
POS Perfume API v1 — Extra Order & Product Endpoints

Routes:
  GET  /api/pos_perfume/v1/orders/stats                     Aggregate order statistics
  GET  /api/pos_perfume/v1/orders/<id>/lines                List lines of a single order
  POST /api/pos_perfume/v1/orders/<id>/lines/bulk           Add multiple lines at once
  GET  /api/pos_perfume/v1/orders/<id>/excel                Download Excel export
  GET  /api/pos_perfume/v1/customers/<id>/orders            Order history for a customer
  GET  /api/pos_perfume/v1/products/<id>/uoms               All UoMs + prices for a product
"""

import io
import logging
from odoo import http
from odoo.http import request, Response

from ._helpers import (
    _success, _error, _check_auth, _serialize_line, _build_line_vals,
)

_logger = logging.getLogger(__name__)


class PosPerfumeApiOrderExtra(http.Controller):

    # -----------------------------------------------------------------------
    # Order statistics
    # -----------------------------------------------------------------------

    @http.route('/api/pos_perfume/v1/orders/stats', type='http', auth='none', methods=['GET'], csrf=False)
    def order_stats(self, date_from='', date_to='', user_id='', partner_id='', **kwargs):
        """Aggregate statistics for POS orders.

        GET /api/pos_perfume/v1/orders/stats
        Query params (all optional):
          date_from   YYYY-MM-DD
          date_to     YYYY-MM-DD
          user_id     filter by salesperson
          partner_id  filter by customer

        Returns:
          total_orders, by_state counts, total_revenue_usd, total_revenue_iqd,
          top_products (top 10 by quantity), top_customers (top 10 by revenue),
          daily_revenue (last 30 days or date range), avg_order_value
        """
        try:
            auth_err = _check_auth()
            if auth_err:
                return auth_err

            # Build base domain
            domain = []
            if date_from:
                domain.append(('date', '>=', date_from))
            if date_to:
                domain.append(('date', '<=', date_to + ' 23:59:59'))
            if user_id:
                domain.append(('user_id', '=', int(user_id)))
            if partner_id:
                domain.append(('partner_id', '=', int(partner_id)))

            # Build extra SQL filters for raw queries
            extra_filters = []
            params: list = []
            if date_from:
                extra_filters.append("o.date >= %s")
                params.append(date_from)
            if date_to:
                extra_filters.append("o.date <= %s")
                params.append(date_to + ' 23:59:59')
            if user_id:
                extra_filters.append("o.user_id = %s")
                params.append(int(user_id))
            if partner_id:
                extra_filters.append("o.partner_id = %s")
                params.append(int(partner_id))

            where_clause = ("AND " + " AND ".join(extra_filters)) if extra_filters else ""

            OrderModel = request.env['pos.perfume.order']
            all_orders = OrderModel.search(domain)

            # --- By-state breakdown ---
            state_counts = {}
            for state_key, _ in OrderModel._fields['state'].selection:
                state_counts[state_key] = 0
            for o in all_orders:
                state_counts[o.state] = state_counts.get(o.state, 0) + 1

            # --- Revenue (only confirmed/done orders) ---
            revenue_domain = domain + [('state', 'in', ['sale', 'done'])]
            revenue_orders = OrderModel.search(revenue_domain)
            total_revenue_usd = sum(o.amount_total for o in revenue_orders)
            total_revenue_iqd = sum(o.amount_total_iqd for o in revenue_orders)
            avg_order_value = (total_revenue_usd / len(revenue_orders)) if revenue_orders else 0.0

            # --- Top 10 products by total quantity (SQL for efficiency) ---
            request.env.cr.execute(f"""
                SELECT
                    p.id            AS product_id,
                    p.name          AS product_name,
                    pt.default_code AS default_code,
                    SUM(l.quantity) AS total_qty,
                    SUM(l.line_total) AS total_revenue
                FROM pos_perfume_order_line l
                JOIN pos_perfume_order o ON o.id = l.order_id
                JOIN product_product pt ON pt.id = l.product_id
                JOIN product_template p ON p.id = pt.product_tmpl_id
                WHERE o.state IN ('sale', 'done')
                {where_clause}
                GROUP BY p.id, p.name, pt.default_code
                ORDER BY total_qty DESC
                LIMIT 10
            """, params)
            top_products = [
                {
                    'product_id':   row[0],
                    'product_name': row[1],
                    'default_code': row[2] or '',
                    'total_qty':    float(row[3] or 0),
                    'total_revenue': float(row[4] or 0),
                }
                for row in request.env.cr.fetchall()
            ]

            # --- Top 10 customers by total revenue (SQL) ---
            request.env.cr.execute(f"""
                SELECT
                    rp.id           AS partner_id,
                    rp.name         AS partner_name,
                    rp.phone        AS phone,
                    COUNT(o.id)     AS order_count,
                    SUM(o.amount_total) AS total_revenue
                FROM pos_perfume_order o
                JOIN res_partner rp ON rp.id = o.partner_id
                WHERE o.state IN ('sale', 'done')
                {where_clause}
                GROUP BY rp.id, rp.name, rp.phone
                ORDER BY total_revenue DESC
                LIMIT 10
            """, params)
            top_customers = [
                {
                    'partner_id':   row[0],
                    'partner_name': row[1],
                    'phone':        row[2] or '',
                    'order_count':  row[3],
                    'total_revenue': float(row[4] or 0),
                }
                for row in request.env.cr.fetchall()
            ]

            # --- Daily revenue for the requested range (or last 30 days) ---
            if date_from or date_to or user_id or partner_id:
                daily_filter  = where_clause
                daily_params  = list(params)
            else:
                daily_filter  = "AND o.date >= NOW() - INTERVAL '30 days'"
                daily_params  = []

            request.env.cr.execute(f"""
                SELECT
                    DATE(o.date)            AS day,
                    COUNT(o.id)             AS orders,
                    SUM(o.amount_total)     AS revenue_usd,
                    SUM(o.amount_total_iqd) AS revenue_iqd
                FROM pos_perfume_order o
                WHERE o.state IN ('sale', 'done')
                {daily_filter}
                GROUP BY DATE(o.date)
                ORDER BY day ASC
            """, daily_params)
            daily_revenue = [
                {
                    'date':        str(row[0]),
                    'orders':      row[1],
                    'revenue_usd': float(row[2] or 0),
                    'revenue_iqd': float(row[3] or 0),
                }
                for row in request.env.cr.fetchall()
            ]

            return _success({
                'total_orders':     len(all_orders),
                'by_state':         state_counts,
                'total_revenue_usd': round(total_revenue_usd, 2),
                'total_revenue_iqd': round(total_revenue_iqd, 2),
                'avg_order_value':  round(avg_order_value, 2),
                'top_products':     top_products,
                'top_customers':    top_customers,
                'daily_revenue':    daily_revenue,
            })
        except Exception as e:
            _logger.error('[POS API] order_stats: %s', e, exc_info=True)
            return _error(str(e), 500)

    # -----------------------------------------------------------------------
    # Order lines — list + bulk add
    # -----------------------------------------------------------------------

    @http.route('/api/pos_perfume/v1/orders/<int:order_id>/lines',
                type='http', auth='none', methods=['GET'], csrf=False)
    def get_order_lines(self, order_id, **kwargs):
        """Get all lines of an order without fetching the full order.

        GET /api/pos_perfume/v1/orders/123/lines
        """
        try:
            auth_err = _check_auth()
            if auth_err:
                return auth_err

            order = request.env['pos.perfume.order'].browse(order_id)
            if not order.exists():
                return _error('Order not found', 404)

            lines = [_serialize_line(l) for l in order.order_line_ids]
            return _success({
                'order_id':    order.id,
                'order_name':  order.name,
                'total_lines': len(lines),
                'lines':       lines,
            })
        except Exception as e:
            _logger.error('[POS API] get_order_lines: %s', e, exc_info=True)
            return _error(str(e), 500)

    @http.route('/api/pos_perfume/v1/orders/<int:order_id>/lines/bulk',
                type='http', auth='none', methods=['POST'], csrf=False)
    def bulk_add_order_lines(self, order_id, **kwargs):
        """Add multiple lines to an order in a single request.

        POST /api/pos_perfume/v1/orders/123/lines/bulk
        Body:
        {
            "lines": [
                { "product_id": 42, "warehouse_id": 1, "quantity": 2.0, "unit_price": 15.0 },
                { "product_id": 55, "warehouse_id": 1, "quantity": 1.0, "unit_price": 25.0, "discount_percent": 10 }
            ]
        }
        Returns: list of created line objects + summary
        """
        try:
            auth_err = _check_auth()
            if auth_err:
                return auth_err

            order = request.env['pos.perfume.order'].browse(order_id)
            if not order.exists():
                return _error('Order not found', 404)
            if order.state in ('done', 'cancel'):
                return _error(f"Cannot add lines to a '{order.state}' order")

            from ._helpers import _get_json_body
            body = _get_json_body()
            lines_data = body.get('lines', [])
            if not lines_data:
                return _error("'lines' array is required and must not be empty")

            created_lines = []
            errors = []

            for idx, line_data in enumerate(lines_data):
                lv = _build_line_vals(line_data)
                if isinstance(lv, str):
                    errors.append({'index': idx, 'error': lv, 'data': line_data})
                    continue
                lv['order_id'] = order.id
                line = request.env['pos.perfume.order.line'].create(lv)
                created_lines.append(_serialize_line(line))

            return _success({
                'order_id':       order.id,
                'created':        len(created_lines),
                'failed':         len(errors),
                'lines':          created_lines,
                'errors':         errors,
                'order_total':    order.amount_total,
                'order_total_iqd': order.amount_total_iqd,
            }, message=f'{len(created_lines)} line(s) added')
        except Exception as e:
            _logger.error('[POS API] bulk_add_order_lines: %s', e, exc_info=True)
            return _error(str(e), 500)

    # -----------------------------------------------------------------------
    # Excel export
    # -----------------------------------------------------------------------

    @http.route('/api/pos_perfume/v1/orders/<int:order_id>/excel',
                type='http', auth='none', methods=['GET'], csrf=False)
    def export_order_excel(self, order_id, **kwargs):
        """Download a formatted Excel invoice for an order.

        GET /api/pos_perfume/v1/orders/123/excel
        Returns the .xlsx binary or JSON error if openpyxl is missing.
        """
        try:
            auth_err = _check_auth()
            if auth_err:
                return auth_err

            order = request.env['pos.perfume.order'].browse(order_id)
            if not order.exists():
                return _error('Order not found', 404)

            try:
                import openpyxl
                from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
                from openpyxl.utils import get_column_letter
            except ImportError:
                return _error('openpyxl is not installed on the server. Run: pip install openpyxl', 500)

            excel_data = _build_excel(order, openpyxl, Font, PatternFill, Alignment, Border, Side, get_column_letter)

            safe_name = (order.name or f'order_{order_id}').replace('/', '-')
            filename = f"Invoice_{safe_name}.xlsx"

            return Response(
                excel_data,
                status=200,
                headers={
                    'Content-Type':        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    'Content-Disposition': f'attachment; filename="{filename}"',
                    'Content-Length':      str(len(excel_data)),
                    'Access-Control-Allow-Origin': '*',
                },
                direct_passthrough=True,
            )
        except Exception as e:
            _logger.error('[POS API] export_order_excel: %s', e, exc_info=True)
            return _error(str(e), 500)

    # -----------------------------------------------------------------------
    # Customer order history
    # -----------------------------------------------------------------------

    @http.route('/api/pos_perfume/v1/customers/<int:partner_id>/orders',
                type='http', auth='none', methods=['GET'], csrf=False)
    def customer_order_history(self, partner_id, state='', limit='50', offset='0',
                               date_from='', date_to='', **kwargs):
        """Paginated order history for a specific customer.

        GET /api/pos_perfume/v1/customers/10/orders?state=sale&limit=20
        Query params:
          state      draft | quotation | sale | done | cancel  (optional)
          limit      default 50
          offset     default 0
          date_from  YYYY-MM-DD
          date_to    YYYY-MM-DD
        """
        try:
            auth_err = _check_auth()
            if auth_err:
                return auth_err

            partner = request.env['res.partner'].browse(partner_id)
            if not partner.exists():
                return _error('Customer not found', 404)

            domain = [('partner_id', '=', partner_id)]
            if state:
                domain.append(('state', '=', state))
            if date_from:
                domain.append(('date', '>=', date_from))
            if date_to:
                domain.append(('date', '<=', date_to + ' 23:59:59'))

            limit_int  = int(limit)
            offset_int = int(offset)

            OrderModel = request.env['pos.perfume.order']
            orders = OrderModel.search(domain, limit=limit_int, offset=offset_int, order='date desc, id desc')
            total  = OrderModel.search_count(domain)

            # Summary stats for this customer
            all_customer_orders = OrderModel.search([
                ('partner_id', '=', partner_id),
                ('state', 'in', ['sale', 'done']),
            ])
            total_spent_usd = sum(o.amount_total for o in all_customer_orders)
            total_spent_iqd = sum(o.amount_total_iqd for o in all_customer_orders)

            items = [
                {
                    'id':              o.id,
                    'name':            o.name,
                    'date':            o.date.isoformat() if o.date else None,
                    'state':           o.state,
                    'invoice_type':    o.invoice_type or '',
                    'amount_total':    o.amount_total,
                    'amount_total_iqd': o.amount_total_iqd,
                    'sap_synced':      o.sap_synced,
                    'sap_doc_num':     o.sap_doc_num or '',
                    'user_name':       o.user_id.name if o.user_id else '',
                    'lines_count':     len(o.order_line_ids),
                }
                for o in orders
            ]

            return _success({
                'partner_id':      partner_id,
                'partner_name':    partner.name,
                'total':           total,
                'offset':          offset_int,
                'limit':           limit_int,
                'summary': {
                    'total_orders':    len(all_customer_orders),
                    'total_spent_usd': round(total_spent_usd, 2),
                    'total_spent_iqd': round(total_spent_iqd, 2),
                },
                'items':           items,
            })
        except Exception as e:
            _logger.error('[POS API] customer_order_history: %s', e, exc_info=True)
            return _error(str(e), 500)

    # -----------------------------------------------------------------------
    # Product UoMs with prices
    # -----------------------------------------------------------------------

    @http.route('/api/pos_perfume/v1/products/<int:product_id>/uoms',
                type='http', auth='none', methods=['GET'], csrf=False)
    def get_product_uoms(self, product_id, pricelist_id='', **kwargs):
        """Get all available UoMs with prices for a product.

        GET /api/pos_perfume/v1/products/42/uoms?pricelist_id=1
        Query params:
          pricelist_id  (optional) — use specific pricelist for prices

        Returns list of { id, name, price } for each available UoM.
        """
        try:
            auth_err = _check_auth()
            if auth_err:
                return auth_err

            product = request.env['product.product'].browse(product_id)
            if not product.exists():
                return _error('Product not found', 404)

            pricelist = None
            if pricelist_id:
                pricelist = request.env['product.pricelist'].browse(int(pricelist_id))
                if not pricelist.exists():
                    pricelist = None

            from odoo.addons.pos_perfume_custom.controllers.pos_perfume_controller import PosPerfumeController
            ctrl = PosPerfumeController()
            uoms = ctrl._get_uoms_from_pricelist(product, pricelist)

            return _success({
                'product_id':   product.id,
                'product_name': product.name,
                'default_code': product.default_code or '',
                'base_uom': {
                    'id':   product.uom_id.id,
                    'name': product.uom_id.name,
                },
                'pricelist_id':  int(pricelist_id) if pricelist_id else None,
                'pricelist_name': pricelist.name if pricelist else None,
                'uoms':          uoms,
            })
        except Exception as e:
            _logger.error('[POS API] get_product_uoms: %s', e, exc_info=True)
            return _error(str(e), 500)


# ---------------------------------------------------------------------------
# Private helper — Excel builder (shared logic extracted from pos_perfume_controller)
# ---------------------------------------------------------------------------

def _build_excel(order, openpyxl, Font, PatternFill, Alignment, Border, Side, get_column_letter):
    """Build and return the raw bytes of an Excel invoice for the given order."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoice"

    # Styles
    header_fill  = PatternFill("solid", fgColor="2563EB")
    section_fill = PatternFill("solid", fgColor="EFF6FF")
    title_fill   = PatternFill("solid", fgColor="1E3A5F")
    white_font   = Font(color="FFFFFF", bold=True, size=12)
    bold_font    = Font(bold=True, size=11)
    title_font   = Font(color="FFFFFF", bold=True, size=14)
    normal_font  = Font(size=11)
    center       = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left         = Alignment(horizontal="left",   vertical="center", wrap_text=True)
    thin         = Side(style="thin", color="BFDBFE")
    thin_border  = Border(left=thin, right=thin, top=thin, bottom=thin)

    def _style(cell, font=None, fill=None, alignment=None, border=None, number_format=None):
        if font:          cell.font          = font
        if fill:          cell.fill          = fill
        if alignment:     cell.alignment     = alignment
        if border:        cell.border        = border
        if number_format: cell.number_format = number_format

    # Column widths: #, Ref, Product, UoM, Warehouse, Qty, Price, Disc%, Total
    for i, w in enumerate([5, 18, 35, 14, 16, 12, 14, 10, 14], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    row = 1

    # Title
    ws.merge_cells(f"A{row}:I{row}")
    _style(ws.cell(row=row, column=1, value="INVOICE / فاتورة"),
           font=title_font, fill=title_fill, alignment=center)
    ws.row_dimensions[row].height = 32
    row += 1

    # Header info
    exchange_rate = float(
        request.env['ir.config_parameter'].sudo().get_param(
            'pos_perfume.default_exchange_rate_usd_iqd', '1550.0'
        )
    )
    state_label = dict(order._fields['state'].selection).get(order.state, order.state)
    header_pairs = [
        ("Invoice No. / رقم الفاتورة", order.name,
         "Date / التاريخ",            order.date.strftime('%Y-%m-%d %H:%M') if order.date else ''),
        ("Customer / العميل",          order.partner_id.name if order.partner_id else '',
         "Phone / الهاتف",            order.partner_id.phone or order.partner_id.mobile or '' if order.partner_id else ''),
        ("Status / الحالة",            state_label,
         "Pricelist / قائمة الأسعار", order.pricelist_id.name if order.pricelist_id else ''),
        ("Salesperson / المندوب",      order.user_id.name if order.user_id else '',
         "Sale Order / أمر البيع",    order.sale_order_id.name if order.sale_order_id else ''),
        ("SAP Doc / مستند SAP",        order.sap_doc_num or '',
         "Exchange Rate / سعر الصرف", f"{exchange_rate:,.0f} IQD/USD"),
    ]
    if order.note:
        header_pairs.append(("Note / ملاحظة", order.note, "", ""))

    for label1, val1, label2, val2 in header_pairs:
        for span, col, val, is_label in [
            ("A:B", 1, label1, True), ("C:E", 3, val1,   False),
            ("F:G", 6, label2, True), ("H:I", 8, val2,   False),
        ]:
            ws.merge_cells(f"{span[0]}{row}:{span[2]}{row}")
            cell = ws.cell(row=row, column=col, value=val)
            _style(cell,
                   font=bold_font if is_label else normal_font,
                   fill=section_fill if is_label else None,
                   alignment=left, border=thin_border)
        ws.row_dimensions[row].height = 18
        row += 1

    row += 1

    # Lines header
    for col_idx, header in enumerate(
        ["#", "Internal Ref / الكود", "Product / المنتج", "UoM / الوحدة",
         "Warehouse / المستودع", "Qty / الكمية", "Unit Price $", "Disc %", "Total $"], 1
    ):
        _style(ws.cell(row=row, column=col_idx, value=header),
               font=white_font, fill=header_fill, alignment=center, border=thin_border)
    ws.row_dimensions[row].height = 22
    row += 1

    # Lines data
    alt_fill  = PatternFill("solid", fgColor="F0F7FF")
    total_usd = 0.0
    for idx, line in enumerate(order.order_line_ids, 1):
        line_fill = alt_fill if idx % 2 == 0 else None
        subtotal  = line.price_subtotal
        total_usd += subtotal
        row_data = [
            idx,
            line.product_id.default_code or '',
            line.product_id.name or '',
            line.product_uom_id.name if line.product_uom_id else '',
            line.warehouse_id.name if line.warehouse_id else '',
            line.quantity,
            line.price_unit,
            line.discount if hasattr(line, 'discount') else 0.0,
            subtotal,
        ]
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row, column=col_idx, value=val)
            _style(cell, font=normal_font, fill=line_fill,
                   alignment=center if col_idx not in (2, 3) else left,
                   border=thin_border,
                   number_format='#,##0.00' if col_idx in (6, 7, 8, 9) else None)
        ws.row_dimensions[row].height = 18
        row += 1

    row += 1

    # Totals
    exchange_rate_val = order.exchange_rate or exchange_rate
    total_iqd = round(total_usd * exchange_rate_val / 1000) * 1000
    for label, value in [
        ("Subtotal / المجموع الجزئي",          f"${order.amount_subtotal:,.2f}"),
        ("Discount / الخصم",                    f"-${order.amount_discount:,.2f}"),
        ("Tax / الضريبة",                       f"${order.amount_tax:,.2f}"),
        ("TOTAL (USD) / المجموع بالدولار",      f"${order.amount_total:,.2f}"),
        ("TOTAL (IQD) / المجموع بالدينار",      f"{int(total_iqd):,} IQD"),
    ]:
        ws.merge_cells(f"A{row}:G{row}")
        ws.merge_cells(f"H{row}:I{row}")
        is_total = "TOTAL" in label
        lbl_font = Font(bold=True, size=12 if is_total else 11, color="FFFFFF" if is_total else "1E3A5F")
        fill     = header_fill if is_total else section_fill
        _style(ws.cell(row=row, column=1, value=label), font=lbl_font, fill=fill, alignment=left,   border=thin_border)
        _style(ws.cell(row=row, column=8, value=value), font=lbl_font, fill=fill, alignment=center, border=thin_border)
        ws.row_dimensions[row].height = 20 if is_total else 18
        row += 1

    # Footer
    row += 1
    ws.merge_cells(f"A{row}:I{row}")
    _style(ws.cell(row=row, column=1, value="Generated by POS Perfume System | نظام نقاط البيع"),
           font=Font(size=9, italic=True, color="6B7280"), alignment=center)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()
